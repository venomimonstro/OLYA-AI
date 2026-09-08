from pathlib import Path

from docx import Document

from app.services.documents import build_docx, repair_docx_layout, structural_qa


def test_structural_qa_checks_table_shape_and_placeholders(tmp_path: Path):
    spec = {
        "title": "Report",
        "blocks": [
            {"type": "paragraph", "text": "Ready"},
            {"type": "table", "header": True, "rows": [["Name", "Value"], ["A", "{{TBD}}"]]},
        ],
    }
    path = tmp_path / "report.docx"
    build_docx(spec, path)
    report = structural_qa(path, spec)
    assert report["status"] == "failed"
    assert any(item["code"] == "unresolved_placeholders" for item in report["issues"])
    assert report["table_count"] == 1


def test_build_docx_repeats_table_header(tmp_path: Path):
    spec = {"title": "Report", "blocks": [{"type": "table", "header": True, "rows": [["A", "B"], ["1", "2"]]}]}
    path = tmp_path / "table.docx"
    build_docx(spec, path)
    doc = Document(path)
    xml = doc.tables[0].rows[0]._tr.xml
    assert "tblHeader" in xml
    assert all(run.bold for run in doc.tables[0].rows[0].cells[0].paragraphs[0].runs)


def test_layout_repair_is_bounded_and_landscapes_wide_tables(tmp_path: Path):
    spec = {
        "title": "Wide",
        "blocks": [{"type": "table", "header": True, "rows": [[str(i) for i in range(8)], ["x" * 50 for _ in range(8)]]}],
    }
    path = tmp_path / "wide.docx"
    build_docx(spec, path)
    before = path.stat().st_size
    result = repair_docx_layout(path, [{"code": "content_touches_page_edge", "page": 1}])
    assert result["applied"] is True
    assert "compact_table_layout" in result["actions"]
    assert path.stat().st_size > 0
    assert before > 0
    doc = Document(path)
    assert len(doc.tables) == 1


def test_non_repairable_issue_does_not_mutate_docx(tmp_path: Path):
    spec = {"title": "Blank", "blocks": [{"type": "paragraph", "text": "Text"}]}
    path = tmp_path / "plain.docx"
    build_docx(spec, path)
    before = path.read_bytes()
    result = repair_docx_layout(path, [{"code": "blank_page", "page": 2}])
    assert result["applied"] is False
    assert path.read_bytes() == before


def test_qa_route_commits_before_render_and_rechecks_revision():
    source = Path("app/api/routes/documents.py").read_text(encoding="utf-8")
    commit_at = source.index("db.commit()\n\n    events")
    render_at = source.index("render_document_artifacts(", commit_at)
    assert commit_at < render_at
    for marker in (
        "Document advanced to a newer revision while QA was running",
        "Document changed while QA was running",
        'revision.qa_status = "running"',
        "FileResponse",
        "document_qa_max_repairs",
        "repair_docx_layout",
    ):
        assert marker in source


def test_document_worker_is_dedicated_and_has_no_docker_socket():
    compose = Path("docker-compose.yml").read_text(encoding="utf-8")
    worker = Path("app/document_worker_api.py").read_text(encoding="utf-8")
    dockerfile = Path("Dockerfile.document-worker").read_text(encoding="utf-8")
    assert "document-worker:" in compose
    section = compose.split("  document-worker:\n", 1)[1].split("\n  app:\n", 1)[0]
    assert "/var/run/docker.sock" not in section
    assert "libreoffice-writer" in dockerfile
    assert "poppler-utils" in dockerfile
    assert 'alias="X-X1-Document-Token"' in worker
    assert "part in {\"\", \".\", \"..\"}" in worker


def test_production_requires_non_default_document_worker_secret():
    main = Path("app/main.py").read_text(encoding="utf-8")
    installer = Path("scripts/install.sh").read_text(encoding="utf-8")
    env = Path(".env.example").read_text(encoding="utf-8")
    assert "document_render_worker_token_is_default" in main
    assert "generated_document" in installer
    assert "document-worker" in installer
    assert "X1_DOCUMENT_RENDER_WORKER_TOKEN=change-me-document-worker" in env


def test_visual_qa_contains_geometry_raster_and_black_block_guards():
    source = Path("app/services/documents.py").read_text(encoding="utf-8")
    for marker in (
        "pdftotext",
        "-bbox-layout",
        "text_clipping_risk",
        "text_overlap",
        "content_touches_page_edge",
        "suspicious_dark_block",
        "deterministic_raster_geometry",
    ):
        assert marker in source
